"""
Document loading & text extraction.
Supports: PDF, DOCX, TXT.
If extraction fails -> fallback descriptor so the UI can show a raw viewer
(QWebEngine native PDF viewer / plain dump) instead of crashing.
"""
from __future__ import annotations

import hashlib
import os
import re
from dataclasses import dataclass, field
from typing import List, Optional


@dataclass
class Section:
    level: int          # 1 = H1, 2 = H2, 3 = H3, 0 = paragraph
    text: str
    sid: str = ""       # stable section id
    kind: str = "p"     # 'p' = paragraph, 'li' = bullet/list item


@dataclass
class LoadedDocument:
    path: str
    title: str
    doc_id: str                      # stable hash id for state keying
    kind: str                        # 'pdf' | 'docx' | 'txt'
    ok: bool                         # extraction succeeded
    fallback_raw: bool = False       # if True, render file natively instead
    sections: List[Section] = field(default_factory=list)
    error: Optional[str] = None

    @property
    def plain_text(self) -> str:
        return "\n".join(s.text for s in self.sections)


def _doc_id_for(path: str) -> str:
    h = hashlib.sha256()
    h.update(os.path.abspath(path).encode("utf-8"))
    try:
        st = os.stat(path)
        h.update(str(st.st_size).encode())
    except OSError:
        pass
    return h.hexdigest()[:16]


# ----------------------------------------------------------------------------
# Heading heuristics
# ----------------------------------------------------------------------------

# Latin digits, Persian digits (۰-۹) and Arabic-Indic digits (٠-٩)
_HEADING_NUM_RE = re.compile(
    r"^\s*((\d|[\u06F0-\u06F9\u0660-\u0669])+([\.\u066B](\d|[\u06F0-\u06F9\u0660-\u0669])+)*"
    r"[\.\)\-\u2013]?|[IVXLC]+\.)\s+\S")
_MD_HEADING_RE = re.compile(r"^(#{1,4})\s+(.*)$")
_BULLET_RE = re.compile(r"^\s*([-*•‣◦▪–]|[\u2022\u25CF\u25AA])\s+(\S.*)$")
# Persian / Arabic structural keywords that begin headings
_FA_HEADING_RE = re.compile(
    r"^\s*(فصل|بخش|قسمت|مبحث|درس|گفتار|مقدمه|نتیجه(\u200cگیری)?|چکیده|خلاصه|"
    r"الفصل|الباب|المبحث|المقدمة|الخاتمة)\b")


def _looks_like_heading(line: str) -> int:
    """Return heading level 1-3 or 0 if not a heading."""
    s = line.strip()
    if not s or len(s) > 90:
        return 0
    m = _MD_HEADING_RE.match(s)
    if m:
        return min(len(m.group(1)), 3)
    # bullets are never headings
    if _BULLET_RE.match(s):
        return 0
    # numbered headings: "1. Intro", "2.3 Methods", "۱. مقدمه", "١٫٢ روش"
    if _HEADING_NUM_RE.match(s) and len(s) < 80 and not s.endswith(('.', '،', '؟', '!', ':', '؛')):
        depth = s.split()[0].count(".") + s.split()[0].count("\u066B") + 1
        return min(max(depth, 1), 3)
    # Persian/Arabic structural keywords: فصل / بخش / مقدمه ...
    if _FA_HEADING_RE.match(s) and len(s) < 70 and not s.endswith(('.', '،', '؟', '!', '؛')):
        return 2
    # ALL CAPS short line (latin)
    letters = [c for c in s if c.isalpha()]
    if letters and len(s) < 60:
        if all(c.upper() == c for c in letters) and len(letters) > 3 and any(c.isupper() for c in letters):
            return 1
    # short line, no terminal punctuation, surrounded by blank lines is decided by caller
    return 0


def _lines_to_sections(lines: List[str]) -> List[Section]:
    sections: List[Section] = []
    buf: List[str] = []

    def flush():
        if buf:
            para = " ".join(x.strip() for x in buf if x.strip())
            if para:
                sections.append(Section(0, para))
            buf.clear()

    prev_blank = True
    for i, raw in enumerate(lines):
        line = raw.rstrip("\n")
        if not line.strip():
            flush()
            prev_blank = True
            continue
        # bullets / list items: keep each as its own block (structure preserved)
        bm = _BULLET_RE.match(line)
        if bm:
            flush()
            sections.append(Section(0, "• " + bm.group(2).strip(), kind="li"))
            prev_blank = False
            continue
        lvl = _looks_like_heading(line)
        # standalone short line between blanks → likely heading
        nxt_blank = (i + 1 >= len(lines)) or (not lines[i + 1].strip())
        if lvl == 0 and prev_blank and nxt_blank and 0 < len(line.strip()) < 70 \
                and not line.strip().endswith(('.', '،', ',', '؛', ';', '?', '؟', '!')):
            lvl = 2
        if lvl > 0:
            flush()
            text = _MD_HEADING_RE.sub(lambda m: m.group(2), line.strip())
            sections.append(Section(lvl, text.strip()))
        else:
            buf.append(line)
        prev_blank = False
    flush()

    # assign ids
    for idx, s in enumerate(sections):
        s.sid = f"sec-{idx}"
    return sections


# ----------------------------------------------------------------------------
# Loaders
# ----------------------------------------------------------------------------

def _load_txt(path: str) -> List[Section]:
    for enc in ("utf-8", "utf-16", "cp1256", "latin-1"):
        try:
            with open(path, "r", encoding=enc) as f:
                return _lines_to_sections(f.readlines())
        except (UnicodeDecodeError, UnicodeError):
            continue
    raise UnicodeError("Could not decode text file")


def _load_pdf(path: str) -> List[Section]:
    from pypdf import PdfReader
    reader = PdfReader(path)
    lines: List[str] = []
    for page in reader.pages:
        txt = page.extract_text() or ""
        lines.extend(txt.splitlines())
        lines.append("")  # page break = blank line
    total = "".join(lines).strip()
    if len(total) < 20:
        raise ValueError("PDF appears to be scanned / empty text layer")
    return _lines_to_sections(lines)


def _load_docx(path: str) -> List[Section]:
    import docx
    d = docx.Document(path)
    sections: List[Section] = []
    for p in d.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "").lower() if p.style else ""
        if "heading" in style:
            m = re.search(r"(\d+)", style)
            lvl = min(int(m.group(1)), 3) if m else 1
            sections.append(Section(lvl, text))
        elif "title" in style:
            sections.append(Section(1, text))
        elif "list" in style or "bullet" in style or "number" in style:
            sections.append(Section(0, "• " + text, kind="li"))
        else:
            sections.append(Section(0, text))
    for idx, s in enumerate(sections):
        s.sid = f"sec-{idx}"
    if not sections:
        raise ValueError("DOCX contained no readable paragraphs")
    return sections


def load_document(path: str) -> LoadedDocument:
    ext = os.path.splitext(path)[1].lower()
    kind = {".pdf": "pdf", ".docx": "docx", ".txt": "txt", ".md": "txt"}.get(ext, "txt")
    title = os.path.splitext(os.path.basename(path))[0]
    doc_id = _doc_id_for(path)
    try:
        if kind == "pdf":
            sections = _load_pdf(path)
        elif kind == "docx":
            sections = _load_docx(path)
        else:
            sections = _load_txt(path)
        return LoadedDocument(path=path, title=title, doc_id=doc_id, kind=kind,
                              ok=True, sections=sections)
    except Exception as e:  # noqa: BLE001 — never crash, fall back
        return LoadedDocument(path=path, title=title, doc_id=doc_id, kind=kind,
                              ok=False, fallback_raw=True, error=str(e))
